#!/usr/bin/env python3
"""Build the protocol / distribution / legacy-freeze reconciliation.

Every legacy (and protocol-listed) analysis item receives exactly one of the
seven allowed labels. The machine-readable table and the Markdown report are
generated together; unknown labels, duplicate IDs or unknown template
placeholders abort the build.

  python3 reconciliation/build_reconciliation.py \
      [--plan PLAN.docx --application APP.docx --dataset DATASET.xlsx \
       --revised-plan REV_PLAN.docx --revised-application REV_APP.docx]

Without the document arguments the stored scope file
(reconciliation/distributed_scope.json) is reused. No ToMMo participant-level
data is read: the dataset workbook is the item dictionary of the planned
distribution, not data.
"""
import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / "reconciliation"
SCOPE = HERE / "distributed_scope.json"
TABLE = HERE / "RECONCILIATION_CLASSIFICATION.csv"
TEMPLATE = HERE / "RECONCILIATION.template.md"
REPORT = ROOT / "PROTOCOL_DISTRIBUTION_FREEZE_RECONCILIATION.md"

C = "KEEP AS CONFIRMATORY"
S = "KEEP AS SECONDARY"
N = "KEEP AS NEGATIVE CONTROL"
E = "KEEP AS EXPLORATORY"
U = "UNTESTABLE WITH DISTRIBUTED DATA"
A = "REQUIRES PROTOCOL AMENDMENT"
X = "SUPERSEDED BY APPROVED-PROTOCOL REQUIREMENT"
LABELS = [C, S, N, E, U, A, X]
GENE_RE = r"\b(TAS1R\d|TAS2R\d+|SCNN1[ABDG]|OTOP1|GNAT3|GNG13|PLCB2|ITPR3|TRPM5|TRPA1)\b"

REV_DRUG = ("06-13修正版が承認版の場合: 同版は5薬剤カテゴリ（Ca拮抗薬/ARB・スタチン・NSAIDs・消化器用薬・"
            "メトホルミン）を『主たる解析対象』と明記し、本予測の薬剤はその列挙外。confirmatory/secondary維持には"
            "同版との整合（列挙外薬剤を事前規定済み仮説として扱えるか）の確認が必要。")
REV_NONE = ""

# (id, legacy source, item, legacy status, label, rationale, note if 2026-06-13 revision governs, final-freeze action)
ITEMS = [
    # ---- directional predictions
    ("P01", "13_v2_FINAL/62/91", "TAS2R38 PAV/AVI diplotype × propylthiouracil (PTU)", "primary (F-PRIMARY)", C,
     "TAS2R38は承認20遺伝子に含まれる。薬剤はDrug ID/KEDD ID→自由記載で同定可能。疎ならn<200で PRE-SPECIFIED BUT UNTESTABLE（73/74/91）。P10による救済禁止。",
     REV_DRUG, "frozen outcome（曝露）とプロトコル目的変数（処方変更）の対応を最終freezeで固定。TAS2R38 3SNPの分譲manifest確認。"),
    ("P02", "13_v2_FINAL/62", "TAS2R38 diplotype × methimazole (thiamazole)", "primary (F-PRIMARY)", C,
     "承認遺伝子・分譲薬剤項目で検証可能。", REV_DRUG, "同上。"),
    ("P03", "13_v2_FINAL/62/72", "TAS2R38 × chloramphenicol (capsule/suspension)", "secondary; 72: PRE-SPECIFIED BUT UNTESTABLE（国内経口使用ほぼ無し）", S,
     "範囲内（承認遺伝子・薬剤項目あり）。検証可否は受領後の genotype-blind 件数で frozen floor により判定。", REV_DRUG,
     "件数floor判定のみ。代替薬への置換禁止。"),
    ("P04", "13_v2_FINAL/62", "TAS2R38 × chlorpheniramine (syrup vs tablet)", "secondary", S, "範囲内。", REV_DRUG, "—"),
    ("P05", "13_v2_FINAL/62/63/64", "TAS2R9 V187A (rs3741845) × ofloxacin", "primary (F-PRIMARY)", C,
     "TAS2R9は承認遺伝子。rs3741845 不在時のproxyはTAS2R9分譲領域内のSNPに限る（64のproxy規則をこの範囲に限定）。",
     REV_DRUG, "分譲SNP manifestで rs3741845 または領域内proxyを確認。"),
    ("P06", "13_v2_FINAL/63", "TAS2R19 R299C (rs10772420) × quinine/quinidine", "secondary; 72: PRE-SPECIFIED BUT UNTESTABLE", A,
     "TAS2R19 は承認20遺伝子に含まれない（申請範囲外）。履歴として保持し、実行しない。", REV_NONE,
     "遺伝子追加の計画変更が承認されない限りUNTESTABLEとして分母外に列挙。"),
    ("P07", "13_v2_FINAL/62", "TAS2R38 × liquid oral medications (EX3 aggregate)", "secondary", S, "範囲内（薬剤項目＋剤形分類）。", REV_DRUG, "—"),
    ("P08", "13_v2_FINAL/63", "TAS1R2 rs12033832 × sweetened formulations", "secondary; 72: CONDITIONALLY TESTABLE", S,
     "TAS1R2は承認遺伝子。甘味剤含有は剤形proxyのみ。", REV_DRUG, "rs12033832 の分譲確認。"),
    ("P09", "13_v2_FINAL/63", "TRPA1 rs11988795 × chemesthetic oral liquids (ibuprofen susp.)", "secondary", A,
     "TRPA1 は承認20遺伝子に含まれない。", REV_NONE, "計画変更なしではUNTESTABLEとして列挙。"),
    ("P10", "13_v2_FINAL/62/91", "TAS2R4 differential-activation variant × thiourea class in AVI carriers", "secondary (62/91); 82・README・LOCK_STATUSはF-PRIMARYと記載", A,
     "TAS2R4 は承認20遺伝子に含まれず、63でも UNRESOLVED。62/91（82 rule 1で正）では secondary。P01救済は禁止のまま。", REV_NONE,
     "F-PRIMARYは P01, P02, P05 と明記（82表・README・LOCK_STATUSの記載矛盾を最終freezeで訂正記録）。"),
    ("P11", "13_v2_FINAL/72", "TAS2R43/TAS2R46 × acesulfame-K sweetened formulations", "secondary; 72: PRE-SPECIFIED BUT UNTESTABLE", U,
     "甘味剤（Ace-K）の同定項目が分譲薬剤項目に無い。TAS2R43も承認外（TAS2R46のみ承認）。", REV_NONE, "UNTESTABLEとして分母外に列挙。"),
    ("P12", "13_v2_FINAL/62", "TAS2R38 × erythromycin suspension vs enteric/film tablet", "secondary", S, "範囲内。", REV_DRUG, "—"),
    # ---- negative controls
    ("N01", "79", "route: TAS2R38 × insulin (parenteral) → null", "control", N, "範囲内。", REV_NONE, "—"),
    ("N02", "79", "route: TAS2R38 × gentamicin (parenteral) → null", "control; 72: UNTESTABLE（外来稀）", N, "範囲内（件数floorで判定）。", REV_NONE, "—"),
    ("N03", "79", "route: TAS2R38 × remdesivir (IV) → null", "control; 72: UNTESTABLE（入院のみ）", N, "範囲内（件数floorで判定）。", REV_NONE, "—"),
    ("N04", "79", "route: TAS2R38 × fentanyl patch (transdermal) → null", "control", N, "範囲内。", REV_NONE, "—"),
    ("N05", "79", "masking: TAS2R38 × omeprazole enteric-coated → null/attenuated", "control", N, "範囲内。", REV_NONE, "—"),
    ("N06", "79", "masking: TAS2R38 × erythromycin enteric/film tablet → attenuated", "control", N, "範囲内（P12と対）。", REV_NONE, "—"),
    ("N07", "79", "receptor-mismatch: TAS2R9 × PROP/thiourea → null", "control", N, "TAS2R9承認。", REV_NONE, "—"),
    ("N08", "79", "receptor-mismatch: TAS2R38 × ofloxacin → null", "control", N, "範囲内。", REV_NONE, "—"),
    ("N09", "79", "receptor-mismatch: TAS2R38 × clindamycin liquid → null", "control; 72: UNTESTABLE", N, "範囲内（件数floorで判定）。", REV_NONE, "—"),
    ("N10", "79", "receptor-mismatch: TAS1R2 × quinine → null", "control; 72: UNTESTABLE", N, "TAS1R2承認。", REV_NONE, "—"),
    ("N11", "79", "masking: TRPA1 × ibuprofen film-coated tablet → attenuated", "control", A, "TRPA1は承認外。", REV_NONE, "計画変更なしではUNTESTABLEとして列挙。"),
    ("N12", "79", "behavioral: TAS2R38 × supplement/health-food use → null", "control", N,
     "服薬項目は『お薬・サプリメント等』を含み構成可能。", REV_NONE, "—"),
    # ---- contrasts
    ("C1", "25/72", "formulation contrast C1", "F-CONTRAST (Holm)", S,
     "範囲内。ただし25（script 09出力: erythromycin suspension vs enteric tab）と72（EX3 vs EX1 within TAS2R38-ligand drugs）で定義が不一致。", REV_DRUG,
     "どちらの定義が凍結対象かを最終freezeで記録（結果を見る前）。"),
    ("C2", "25/72", "formulation/route contrast C2", "F-CONTRAST", S,
     "25: acetaminophen liquid vs film tablet（根拠受容体TAS2R39は承認外）／72: oral liquid vs parenteral same drug。定義不一致・遺伝子割当未規定。", REV_DRUG,
     "遺伝子割当がTAS2R39を要する場合はREQUIRES PROTOCOL AMENDMENTへ移す旨を最終freezeで判断。"),
    ("C3", "25/72", "formulation contrast C3", "F-CONTRAST", S, "25: theophylline elixir vs SR tablet／72: erythromycin suspension vs tablet。定義不一致。", REV_DRUG, "同上。"),
    ("C4", "25/72", "formulation contrast C4", "F-CONTRAST", S, "25: KCl liquid vs ER tablet／72: OD vs standard tablet。定義不一致。", REV_DRUG, "同上。"),
    ("C5", "25/72", "route/release contrast C5", "F-CONTRAST", S, "25: insulin route contrast／72: enteric vs non-enteric same ingredient。定義不一致。", REV_NONE, "同上。"),
    # ---- gradient / calibration / rubric / hierarchy
    ("GR1", "95/83§4-G/80", "formal 1-df genotype × ordinal exposure (EX0–EX3) interaction β3", "F-GRADIENT, principal mechanistic test", C,
     "分譲データ（薬剤項目＋剤形）で構成可能。G（どの遺伝子型）とY（どの薬剤集合）が95で一意に規定されていない。", REV_NONE,
     "G/Y の一意化と、共変量からancestry PCを除く扱い（GR/SAP-COV2参照）を最終freezeで固定。parser実文字列QC合格が前提。"),
    ("GR2", "95/80", "categorical-E secondary model and order-restricted |β_EX3|≥…≥|β_EX0|", "secondary within F-GRADIENT", S, "範囲内。", REV_NONE, "—"),
    ("GR3", "80", "descriptive ordered-gradient table (G1/S1/G→E/G→F/G→T/control)", "descriptive", S,
     "S1（知覚）は UNMEASURED のまま。", REV_NONE, "—"),
    ("CA1", "96/81", "continuous calibration: Spearman(evidence score, signed Z_i), ≥10,000 permutations, seed 42", "F-CALIBRATION", C,
     "範囲内。ただしP06/P09/P10/P11が分母外となるため評価可能点は最大8（P01–P05,P07,P08,P12）で、受領前にpowerが低いことを明記。", REV_NONE,
     "最小評価点数（UNDEFINED扱いの閾値）を最終freezeで明記。"),
    ("CA2", "32/81", "binary supported/not-supported calibration", "secondary within F-CALIBRATION (82 rule 4)", S, "範囲内。", REV_NONE, "—"),
    ("RB1", "21/23", "evidence-score rubric G+H+R+V+Rep+F+O−A and frozen scores P01–P12", "frozen input to calibration", C,
     "DRY_REPRODUCTION_AUDITで全12点を再計算し一致。", REV_NONE, "変更不可。"),
    ("HI1", "62 (+82 rule 1)", "prediction hierarchy: primary = P01, P02, P05; others secondary; N01–N12 control", "frozen", C,
     "62が正（82 rule 1）。82表・README_HANDOFF・LOCK_STATUSの『F-PRIMARY = P01, P02, P05, P10』は62と矛盾する記載誤り。", REV_NONE,
     "最終freezeで訂正をerrataとして記録（履歴ファイルは書換えない）。"),
    # ---- SAP
    ("SAP-POP", "83§1", "analysis population: genotyped, wave-1 medication answered; wave-2 analyses in both-wave subset", "frozen", X,
     "承認計画は登録時20歳以上と重大欠測除外を要求（旧定義に年齢条件なし）。", "06-13修正版ではさらにベースライン・第2段階の両方参加者に限定。",
     "20歳以上（＋承認版により両wave参加）を最終freezeに記載。"),
    ("SAP-G", "83§2/63/65", "genotype coding: TAS2R38 PAV dosage, other effect-allele dosage", "frozen", C, "承認遺伝子に限り適用。", REV_NONE, "—"),
    ("SAP-O1", "83§3", "O1 exposure presence (drug/class), binary", "primary outcome", S,
     "承認計画の目的変数は処方変更（用量変更・薬剤変更: switch/addition/deletion）であり、現在使用は主要目的変数になり得ない。O1自体はwave別横断指標として保持。", REV_NONE,
     "P01/P02/P05 の主要検定をどの目的変数で行うかを最終freezeで固定。"),
    ("SAP-O2", "83§3", "O2 EX3 vs EX1 within drug", "primary-family outcome", S, "範囲内（剤形）。", REV_NONE, "—"),
    ("SAP-O3", "83§3/74/78", "O3 cross-wave (ph1→ph2) reported-exposure transition/discordance", "secondary", C,
     "承認計画の目的変数（処方変更）に対応するため主要へ格上げ。用語は78に従い『reported exposure transition』（adherence/意図/味覚駆動とは呼ばない）。", REV_NONE,
     "用量変更・有効成分変更・同一治療目的内switch・追加・削除の操作的定義を最終freezeで規定。"),
    ("SAP-O4", "83§3", "O4 EX3 share (ordinal/continuous)", "frozen", S, "範囲内。", REV_NONE, "—"),
    ("SAP-M1", "83§4", "O1/O2 logistic; O3 logistic with participant-clustered SE; O4 linear/ordinal; OR+95%CI; sign check before p", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("SAP-COV1", "83§4/95", "legacy covariate set: age + sex + wave + 10 ancestry PCs (+indication proxy)", "frozen", X,
     "承認計画は年齢・性別・既往歴・併存疾患・生活習慣・社会経済指標の調整を要求。", REV_NONE, "共変量集合を分譲項目名で最終freeze。"),
    ("SAP-COV2", "83§4/95", "10 genome-wide ancestry principal components", "frozen covariate", U,
     "分譲は候補遺伝子領域SNPのみでPC算出不可、PC変数もデータセット仕様に無い。", REV_NONE, "代替（例: 出生地コード）を採る場合は最終freezeで明記。"),
    ("SAP-QC1", "83§2/87", "genotype QC: MAF, HWE, missingness (genotype-blind)", "frozen", C, "範囲内（分譲SNPに対して）。", REV_NONE, "—"),
    ("SAP-QC2", "83§2/87", "genotype QC: relatedness and ancestry PCs", "frozen", U,
     "genome-wide SNP・血縁ID がデータセット仕様に無い。", REV_NONE, "—"),
    ("SAP-EST", "83§5", "estimand: population-average genotype-associated difference in past-two-week reported exposure; not adherence/perception", "frozen", C, "範囲内。", REV_NONE, "処方変更目的変数への拡張時も同じ否定条件を維持。"),
    ("SAP-MISS", "83§6", "complete-case; MI only as sensitivity", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("SAP-ORDER", "83§7", "ordered execution: genotype-blind QC → analysis-set lock → primary → gradient → controls → contrasts → calibration → 85", "frozen", C,
     "範囲内。前段に final pre-receipt freeze → independent timestamp → receipt を追加。", REV_NONE, "—"),
    ("SAP-SW", "83§8", "software freeze", "frozen", C, "範囲内。", REV_NONE, "—"),
    # ---- sensitivity
    ("S1", "84", "unspecified-tablet EX2→EX1 reclassification", "sensitivity", S, "範囲内。", REV_NONE, "—"),
    ("S2", "84/64", "P05 proxy-variant swap", "sensitivity", S, "proxyはTAS2R9分譲領域内に限る。", REV_NONE, "—"),
    ("S3", "84/65", "A49P-only TAS2R38 fallback", "sensitivity", S, "範囲内。", REV_NONE, "—"),
    ("S4", "84", "age ≥40 subset", "sensitivity", S, "範囲内。", REV_NONE, "—"),
    ("S5", "84", "prescription-only exposures (入手方法=処方箋)", "sensitivity", S, "入手方法項目が分譲仕様に存在。", REV_NONE, "—"),
    ("S6", "84/75", "claims-defined exposure replication", "sensitivity (conditional)", U,
     "承認計画は医科レセプト・介護保険の利用に言及するが、分譲予定データセット仕様に含まれない。", REV_NONE, "受領manifestで確認されるまでNOT RUN。"),
    ("S7", "84", "TAS2R38 allelic vs diplotype coding", "sensitivity", S, "範囲内。", REV_NONE, "—"),
    ("S8", "84", "add smoking + alcohol to models", "sensitivity", X, "承認計画は生活習慣の調整を主解析で要求するため、感度解析から主モデル共変量へ移る。", REV_NONE, "—"),
    ("S9", "84", "E-value", "sensitivity", S, "範囲内。", REV_NONE, "—"),
    ("S10", "84", "multiple imputation (MAR)", "sensitivity", S, "範囲内。", REV_NONE, "—"),
    # ---- multiplicity
    ("MT-P", "82", "F-PRIMARY Holm α=0.05", "frozen", C, "構成は62に従い P01, P02, P05。", REV_NONE, "—"),
    ("MT-S", "82", "F-SECONDARY BH q=0.10", "frozen", S, "分母は承認範囲内の secondary 予測のみ（P06/P09/P10/P11は分母外として列挙）。", REV_NONE, "—"),
    ("MT-N", "82", "F-CONTROL descriptive", "frozen", N, "N11は分母外として列挙。", REV_NONE, "—"),
    ("MT-C", "82", "F-CONTRAST Holm α=0.05", "frozen", S, "C1–C5定義の確定が前提。", REV_NONE, "—"),
    ("MT-G", "82/95", "F-GRADIENT single test α=0.05", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("MT-CAL", "82/96", "F-CALIBRATION single test α=0.05", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("MT-X", "82", "F-EXPLORATORY unadjusted, labelled", "frozen", E, "範囲内。", REV_NONE, "—"),
    ("MT-R2", "82 rule 2", "UNDEFINED tests removed from denominator before correction and listed", "frozen", C, "範囲内。", REV_NONE, "—"),
    # ---- parser / exposure classification
    ("PA1", "71/94", "medication parser v1.0.0 (SHA-256 c722341a…) for formulation/route/EX class", "frozen", C, "範囲内。再実行でgold 146件の出力が凍結結果と一致。", REV_NONE, "—"),
    ("PA2", "72", "free text as the primary basis of active-ingredient identification", "frozen premise", X,
     "分譲仕様に KEDD ID・Drug ID（45 slot×ph1/ph2）があるため、成分同定は structured ID 優先、自由記載はbrand/剤形/経路補完とQCへ（本セッション指示）。", REV_NONE,
     "ID→成分→同一治療目的クラスの対応表を最終freezeで凍結。"),
    ("PA3", "92/87/98-R1", "real ToMMo-string parser QC gate (300–500 genotype-blind strings) before any association", "required gate", C, "範囲内。受領後・genotype非結合で実施。", REV_NONE, "—"),
    ("PA4", "69", "formulary gold-corpus validation 145/146 (known failure: オオサカ堂 ビタミンC)", "frozen", C, "再計算で一致。", REV_NONE, "—"),
    ("PA5", "66/67/68", "formulation dictionary, EX0–EX3 oral-sensory exposure classes, taste-masking table", "frozen", C, "範囲内。", REV_NONE, "—"),
    # ---- genotype implementation
    ("GT1", "63/65", "TAS2R38 rs713598/rs1726866/rs10246939 → PAV/AVI haplotype", "READY_WITH_TRANSFORMATION", C, "承認遺伝子。", REV_NONE, "分譲manifest確認。"),
    ("GT2", "63", "TAS2R19 rs10772420", "READY", A, "承認外遺伝子。", REV_NONE, "—"),
    ("GT3", "63", "TRPA1 rs11988795", "READY", A, "承認外遺伝子。", REV_NONE, "—"),
    ("GT4", "63", "TAS1R2 rs12033832", "READY", C, "承認遺伝子（P08/N10で使用）。", REV_NONE, "分譲manifest確認。"),
    ("GT5", "63/64", "TAS2R9 rs3741845 (+pre-specified proxy)", "READY_WITH_PRESPECIFIED_PROXY", C, "承認遺伝子。proxyは領域内に限る。", REV_NONE, "分譲manifest確認。"),
    ("GT6", "63", "TAS2R4 differential-activation variant", "UNRESOLVED", A, "承認外遺伝子かつ変異未特定。", REV_NONE, "—"),
    # ---- waves / power / confounding / terminology / interpretation / governance
    ("RW1", "74 rule 1", "unit = participant × wave exposure state", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("RW2", "74 rule 2/78", "two-week instrument cannot measure adherence/persistence/discontinuation/switching intent", "frozen", C, "範囲内（処方変更目的変数にも適用）。", REV_NONE, "—"),
    ("RW3", "74 rule 3", "cross-wave analyses secondary/supportive, never primary", "frozen", X,
     "承認計画の目的変数はph1→ph2の処方変更であり、縦断変化を主要とする必要がある。", "06-13修正版は両wave参加者限定を明記（同趣旨）。", "—"),
    ("RW4", "74 rule 4/73/91", "sparse-cell floor n<200 → PRE-SPECIFIED BUT UNTESTABLE; no merging/rescue", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("RW5", "74 rule 5", "repeated measures with GEE / clustered SE", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("RW6", "74 rule 6/76/77", "wave-varying indication confounding; DAG-based indication proxies", "frozen", C,
     "範囲内（疾患・治療関連の調査票項目が分譲仕様にある）。", REV_NONE, "—"),
    ("PW1", "73/10_power_assessment.py", "pre-analysis power assessment (seed 42; rare drug × rare variant MARGINAL)", "frozen", S,
     "再実行で9シナリオ全一致。floor値はRW4で拘束。", "06-13修正版の想定標本数（両時点服薬記録 約20,000–25,000）は旧前提と別に記録。", "—"),
    ("TM1", "78", "outcome terminology freeze (reported exposure; banned: adherence, persistence, taste-driven, preference)", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("IN1", "85", "interpretation matrix A–E with pre-registered assignment order", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("IN2", "16/75/85 rule 4", "claim ceiling Level 4; Level-5 causal claims banned", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("IN3", "86", "journal decision tree", "frozen", S, "出版判断であり検定ではない。", REV_NONE, "—"),
    ("AU1", "87", "analysis authorization matrix (allowed/conditional/forbidden operations)", "frozen", C, "範囲内。", REV_NONE, "—"),
    ("AU2", "87/89/98", "'explicit OPEN decision (89/98)' as trigger for genotype–outcome operations", "frozen trigger", X,
     "本セッションの順序（reconciliation → final pre-receipt freeze → independent timestamp → receipt → offline validation）と承認範囲に置換。歴史的OPENは参加者レベル解析の許可にならない。", REV_NONE, "—"),
    ("AM1", "88", "amendment policy", "frozen", C, "範囲内。", REV_NONE, "—"),
    # ---- gates and errata
    ("GA1", "docs/17", "Stage 0–2 GO/NO-GO: CONDITIONAL GO", "historical", S, "再現可能な定量入力を監査（DRY_REPRODUCTION_AUDIT）。条件は後続lockに継承。", REV_NONE, "—"),
    ("GA2", "36", "feasibility gate: CONDITIONAL GO TO PRIMARY VALIDATION", "historical", S,
     "件数記載の誤り（8→9）は61で訂正済み。条件1の遺伝子リストにTAS2R19/TRPA1を含む（承認外）。", REV_NONE, "—"),
    ("GA3", "61", "pre-opening errata E1 (9 predictions with score ≥9)", "errata", C, "再計算で9件（P01–P07, P09, P12）を確認。", REV_NONE, "—"),
    ("GA4", "89", "final opening decision v1: OPEN WITH RESTRICTIONS", "superseded by 98", X, "98で置換済み、かつAU2に従い認可効力なし。", REV_NONE, "—"),
    ("GA5", "98", "final opening decision v2: OPEN PRIMARY DATA WITH RESTRICTIONS", "historical gate", X,
     "historical gate decision。本セッションでは参加者レベル解析を開始する根拠にしない（AU2）。", REV_NONE, "—"),
    ("GA6", "98 restrictions 1–6", "binding restrictions: parser gate, variant availability/floors, P01/P10 separation, denominator reporting, instrument/claims limits, no isolated-association framing", "binding", C, "範囲内。", REV_NONE, "—"),
    # ---- resources and protocol-listed analyses
    ("RS1", "75", "claims (医科レセプト) / dispensing linkage module", "secondary optional", U,
     "承認計画の必要性欄は医科レセプト・介護保険に言及するが、分譲予定データセット仕様（xlsx）に無い。", REV_NONE, "受領manifestで確認されるまで主解析の前提にしない。"),
    ("RS2", "75/80", "G → S1 direct sensory (psychophysical) phenotype", "banned (no taste phenotype assumed)", U,
     "心理物理学的味覚測定は分譲に無い。ただし分譲仕様には自己申告の味覚項目（ph1: 甘・塩・酸・苦・うま味を感じない、常に口の中が苦い等／ph2: 味を感じにくい 1項目）があり、旧freezeの『味覚表現型なし』という前提は事実として不正確。", REV_NONE,
     "自己申告味覚項目を使うかは最終freezeで事前規定するまで解析しない（本文書では新規解析を追加しない）。"),
    ("RS3", "docs/17 condition 5", "pediatric-style rejection endpoints", "conditional", U, "解析対象は20歳以上。", REV_NONE, "—"),
    ("RS4", "75/84 S8/docs/27", "FFQ ph1/ph2 (coffee/tea/alcohol etc.) use", "covariate / sensitivity only", S,
     "分譲仕様にFFQ ph1/ph2あり。旧freezeに変数レベルのtriangulation仕様は無い。", REV_NONE,
     "限定したchemosensory phenotypeのsecondary triangulationとして変数・方向・検定を最終freezeで規定。"),
    ("RS5", "keikaku 統計解析", "regularized regression / gradient boosting with internal cross-validation", "not in legacy freeze (protocol)", E,
     "承認計画に記載。旧freezeの検定階層外なので確認的主張には使わない。", REV_NONE, "—"),
    ("RS6", "keikaku 統計解析", "external validation", "not in legacy freeze (protocol)", U, "外部コホートは分譲に無い。", REV_NONE, "—"),
    ("RS7", "keikaku 統計解析", "survival analysis of prescription change", "not in legacy freeze (protocol)", U,
     "2時点の横断調査で日付付き処方イベントが無い。", REV_NONE, "—"),
    ("RS8", "keikaku 統計解析", "genome-wide association (GWAS)", "not in legacy freeze (protocol, exploratory)", U,
     "分譲は候補遺伝子領域SNPのみ。", "06-13修正版はGWASを候補遺伝子領域解析に置換済み（整合）。", "—"),
    ("RS9", "keikaku 統計解析", "candidate-region association across the 20 approved genes", "not in legacy freeze (protocol)", E,
     "承認範囲内。旧予測にない遺伝子・SNPはF-EXPLORATORY。", REV_NONE, "—"),
    ("RS10", "keikaku 目的変数", "tolerability auxiliary indicators (adverse events, laboratory values)", "not in legacy freeze (protocol)", E,
     "laboratory_test ph1/ph2は分譲仕様にある。有害事象項目は未確認。", REV_NONE, "—"),
    ("RS11", "keikaku 主研究への統合", "providing genotype–prescription-change estimates to the main study (immune profile / clinical outcome)", "out of legacy scope", U,
     "免疫・臨床転帰データはToMMo分譲に無い。本研究は推定値の提供まで。", REV_NONE, "—"),
    ("RS12", "keikaku 情報", "metabolomics", "not used", U, "分譲仕様に含まれない。", REV_NONE, "—"),
    # ---- research-arc future hypotheses
    ("FH1", "research arc", "acquired sensory plasticity of chemosensory phenotype", "future hypothesis", U, "ToMMo分譲で検証不能。", REV_NONE, "—"),
    ("FH2", "research arc", "食育 / sensory conditioning intervention", "future hypothesis", U, "介入データ無し。", REV_NONE, "—"),
    ("FH3", "research arc", "modification of medication response by adjusting sensory phenotype", "future hypothesis", U, "介入データ無し。", REV_NONE, "—"),
    ("FH4", "research arc", "health-economic benefit (wider use of inexpensive/off-patent drugs)", "future hypothesis", U, "費用データ無し。", REV_NONE, "—"),
    ("FH5", "research arc", "global-health benefit", "future hypothesis", U, "検証不能。", REV_NONE, "—"),
    ("FH6", "research arc / keikaku", "immune mediation of chemosensory–drug association", "future hypothesis", U, "免疫データ無し。", REV_NONE, "—"),
]


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def docx_text(p):
    import docx
    d = docx.Document(p)
    out = [x.text for x in d.paragraphs if x.text.strip()]
    for t in d.tables:
        for r in t.rows:
            out.extend(c.text for c in r.cells)
    return "\n".join(out)


def build_scope(a):
    scope = {"documents": {}}
    for key in ["plan", "application", "dataset", "revised_plan", "revised_application"]:
        p = getattr(a, key)
        if p:
            scope["documents"][key] = {"file": Path(p).name, "sha256": sha(p), "bytes": Path(p).stat().st_size}
    plan = docx_text(a.plan)
    line = next(l for l in plan.splitlines() if "さらに具体的な遺伝子名称" in l)
    scope["plan_gene_line_genes"] = sorted(set(re.findall(GENE_RE, line)))
    scope["plan_gene_line_typo"] = "AS2R8" in line and "TAS2R8" not in line
    app = docx_text(a.application)
    scope["application_gene_appendix"] = sorted(set(re.findall(GENE_RE, app)))
    scope["plan_mentions_claims"] = "医科レセプト" in plan
    scope["plan_mentions_gwas"] = "GWAS" in plan
    scope["plan_age20"] = "20歳以上" in plan
    scope["plan_both_waves"] = "双方に参加" in plan
    if a.revised_plan:
        rp = docx_text(a.revised_plan)
        scope["revised_plan_both_waves"] = "双方に参加" in rp
        scope["revised_plan_candidate_categories"] = "候補薬剤カテゴリ" in rp
        scope["revised_plan_gwas_replaced"] = "候補領域に絞った関連解析" in rp
    if a.revised_application:
        scope["revised_application_gene_appendix"] = sorted(set(re.findall(GENE_RE, docx_text(a.revised_application))))
    xl = pd.read_excel(a.dataset, sheet_name=None, header=None)
    rows = []
    for name, df in xl.items():
        for v in df.astype(str).values.tolist():
            rows.append((name, " ".join(x for x in v if x != "nan")))
    txt = [r[1] for r in rows]
    def cnt(*kw):
        return sum(any(k in t for k in kw) for t in txt)
    scope["dataset"] = {
        "sheets": list(xl), "nonempty_rows": len(txt),
        "kedd_id_rows": cnt("KEDD ID"), "drug_id_rows": cnt("Drug ID"),
        "medication_name_rows": cnt("商品名や成分等"), "prescription_source_rows": cnt("医師の処方箋"),
        "dose_rows": cnt("1回の使用量", "１回の使用量"), "frequency_rows": cnt("使用頻度"),
        "duration_rows": cnt("使用期間"), "taste_rows": cnt("味覚", "味を感じ"),
        "ffq_rows": cnt("qa_ffq", "食物摂取頻度", "FFQ", "調査票（食）"), "claims_rows": cnt("レセプト"),
        "metabolome_rows": cnt("メタボローム"), "ancestry_pc_rows": cnt("主成分"),
        "genotype_rows": cnt("SNP", "遺伝子型", "ジェノタイプ"),
    }
    sel = pd.read_excel(a.dataset, sheet_name="分譲対象", header=None)
    files = []
    for _, v in sel.iloc[5:].iterrows():
        vals = [x for x in v.tolist() if str(x) != "nan"]
        flags = [x for x in vals if isinstance(x, (int, float)) and x in (0, 1)]
        if flags:
            label = " / ".join(str(x) for x in vals if not isinstance(x, (int, float)) and not str(x).startswith("(1="))
            files.append({"item": label, "selected": int(flags[0])})
    scope["dataset"]["selected_files"] = [f["item"] for f in files if f["selected"]]
    scope["dataset"]["not_selected"] = [f["item"] for f in files if not f["selected"]]
    scope["dataset"]["release"] = str(sel.iloc[0, 3])
    # approved universe = attached plan gene line (+ attached application); the revised
    # application is recorded for comparison only and never merged in.
    approved = set(scope["plan_gene_line_genes"]) | set(scope["application_gene_appendix"])
    if scope["plan_gene_line_typo"]:
        approved.add("TAS2R8")
    scope["approved_genes"] = sorted(approved)
    scope["approved_genes_source"] = ("attached plan gene line" + (" + TAS2R8 restored from typo 'AS2R8'" if scope["plan_gene_line_typo"] else ""))
    rev = set(scope.get("revised_application_gene_appendix", []))
    scope["revised_vs_approved_genes"] = {"only_revised": sorted(rev - approved), "only_attached": sorted(approved - rev)} if rev else {}
    SCOPE.write_text(json.dumps(scope, indent=1, ensure_ascii=False, sort_keys=True))
    return scope


def main():
    ap = argparse.ArgumentParser()
    for k in ["plan", "application", "dataset", "revised_plan", "revised_application"]:
        ap.add_argument("--" + k.replace("_", "-"), dest=k)
    a = ap.parse_args()
    scope = build_scope(a) if a.plan else json.loads(SCOPE.read_text())
    ids = [r[0] for r in ITEMS]
    dup = [k for k, v in Counter(ids).items() if v > 1]
    bad = [r[0] for r in ITEMS if r[4] not in LABELS]
    if dup or bad:
        sys.exit(f"duplicate ids {dup}; invalid labels {bad}")
    cols = ["item_id", "legacy_source", "item", "legacy_status", "classification", "rationale",
            "if_2026-06-13_revision_governs", "final_pre_receipt_freeze_action"]
    df = pd.DataFrame(ITEMS, columns=cols)
    df.to_csv(TABLE, index=False)
    # legacy genes vs approved universe (computed, not typed)
    v2 = pd.read_csv(ROOT / "13_LOCKED_PREDICTIONS_v2_FINAL.csv")
    legacy = sorted({g for s in v2.gene.dropna() for g in re.findall(GENE_RE + r"|TAS2R\d+", s)} |
                    {g for s in v2.gene.dropna() for g in re.split(r"[/ ]", s) if g})
    approved = set(scope["approved_genes"])
    gene_rows = ["| legacy gene | in approved 20-gene list | used by |", "|---|---|---|"]
    for g in sorted(set(legacy)):
        used = ", ".join(v2.prediction_id[v2.gene.fillna("").str.contains(rf"\b{g}\b", regex=True)])
        gene_rows.append(f"| {g} | {'yes' if g in approved else '**no**'} | {used} |")
    esc = lambda s: str(s).replace("|", "/").replace("\n", " ")
    tbl = ["| ID | source | item | legacy status | **classification** | rationale | if 06-13 revision governs | final-freeze action |",
           "|---|---|---|---|---|---|---|---|"]
    for r in ITEMS:
        tbl.append("| " + " | ".join(esc(x) if i != 4 else f"**{x}**" for i, x in enumerate(r)) + " |")
    counts = Counter(r[4] for r in ITEMS)
    lc = ["| label | n |", "|---|---|"] + [f"| {l} | {counts.get(l, 0)} |" for l in LABELS] + [f"| total | {len(ITEMS)} |"]
    d = scope["dataset"]
    docs = ["| role | file | bytes | SHA-256 |", "|---|---|---|---|"] + [
        f"| {k} | {v['file']} | {v['bytes']} | `{v['sha256']}` |" for k, v in scope["documents"].items()]
    fill = {
        "CLASSIFICATION_TABLE": "\n".join(tbl), "LABEL_COUNTS": "\n".join(lc), "GENE_TABLE": "\n".join(gene_rows),
        "DOC_TABLE": "\n".join(docs), "N_ITEMS": len(ITEMS),
        "APPROVED_GENES": ", ".join(scope["approved_genes"]), "N_APPROVED": len(scope["approved_genes"]),
        "PLAN_GENE_LINE_GENES": ", ".join(scope["plan_gene_line_genes"]),
        "PLAN_GENE_TYPO": "yes（本文に『TTAS2R群（AS2R8, …』と記載、TAS2R8 が誤記）" if scope["plan_gene_line_typo"] else "no",
        "PLAN_CLAIMS": "yes" if scope["plan_mentions_claims"] else "no",
        "PLAN_BOTH_WAVES": "yes" if scope["plan_both_waves"] else "no",
        "REV_BOTH_WAVES": "yes" if scope.get("revised_plan_both_waves") else "no",
        "REV_CATEGORIES": "yes" if scope.get("revised_plan_candidate_categories") else "no",
        "REV_GWAS": "yes" if scope.get("revised_plan_gwas_replaced") else "no",
        **{f"DS_{k.upper()}": v for k, v in d.items() if not isinstance(v, list)},
        "DS_SHEETS": ", ".join(d["sheets"]),
        "DS_SELECTED": "; ".join(d["selected_files"]), "DS_NOT_SELECTED": "; ".join(d["not_selected"]),
        "APPROVED_SOURCE": scope["approved_genes_source"],
        "REV_GENE_DIFF": (json.dumps(scope["revised_vs_approved_genes"], ensure_ascii=False) if scope.get("revised_vs_approved_genes") else "n/a"),
        **{f"N_{l.split()[-1].upper() if l != X else 'SUPERSEDED'}": counts.get(l, 0) for l in LABELS},
    }
    tpl = TEMPLATE.read_text()
    keys = set(re.findall(r"\{\{(\w+)\}\}", tpl))
    miss = sorted(keys - set(fill))
    if miss:
        sys.exit(f"unknown placeholders {miss}")
    for k in keys:
        tpl = tpl.replace("{{" + k + "}}", str(fill[k]))
    REPORT.write_text(tpl)
    print(f"{len(ITEMS)} items -> {TABLE.name}, {REPORT.name}; " + "; ".join(f"{l}={counts.get(l,0)}" for l in LABELS))


if __name__ == "__main__":
    main()
